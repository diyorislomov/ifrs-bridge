import re

from pdfminer.high_level import extract_text as _pdfminer_extract_text
from openpyxl import load_workbook

# An amount token: optional leading minus/parenthesis, digits grouped by
# comma/space/dot separators, optional decimal part, optional closing
# parenthesis (accounting notation for negatives).
_AMOUNT_TOKEN = r'\(?-?\d[\d,.\s\xa0]{0,20}\d\)?'
# A trailing amount at the end of a line.
_AMOUNT_RE = re.compile(_AMOUNT_TOKEN + r'\s*$')
# Same as _AMOUNT_TOKEN but WITHOUT internal spaces, used only for the
# two-column pattern below. Two separate space-grouped amounts on one line
# (e.g. "30 000 000  25 000 000") are ambiguous with a single space-grouped
# amount ("30 000 000") -- allowing spaces inside each column's token there
# would let a single number get mis-split across the two columns. Comma/dot
# grouping doesn't have this ambiguity, so it's still supported.
_AMOUNT_TOKEN_NO_SPACE = r'\(?-?\d[\d,.]{0,20}\d\)?'
# Two trailing amounts (current + prior period columns), with a label
# before them that contains at least one letter -- [^\W\d_] is a
# Unicode-aware "is a letter" check, so this covers Latin and Cyrillic
# (including Ў Қ Ғ Ҳ) without hardcoding specific alphabets.
_TWO_COLUMN_RE = re.compile(
    rf'^(?P<label>.*?[^\W\d_].*?)\s+(?P<cur>{_AMOUNT_TOKEN_NO_SPACE})\s+(?P<prev>{_AMOUNT_TOKEN_NO_SPACE})\s*$'
)
# An optional numeric account code at the start of a line (e.g. "1000",
# "10.1"), separated from the account name by whitespace.
_LEADING_CODE_RE = re.compile(r'^\s*(\d{1,6}(?:[.\-]\d+)?)\s+')
# Date header lines like "June 30, 2026", which would otherwise look like
# a label followed by a trailing amount ("2026").
_DATE_LINE_RE = re.compile(r',\s*(19|20)\d{2}\s*$')
# Company legal-entity suffix (Latin and Cyrillic, Uzbek and international).
_COMPANY_SUFFIX_RE = re.compile(r'\b(?:Ltd|LLC|JSC|MChJ|AJ|АЖ|МЧЖ|QK|ҚК)\b', re.IGNORECASE)
# A quoted company name, e.g. '"Example"' or '«Namuna»'.
_QUOTED_NAME_RE = re.compile(r'[«"]([^«»"]{2,80})[»"]')

# Section headers (Uzbek Cyrillic + English), used to split the document
# into rough balance-sheet / P&L / cash-flow / notes regions for a more
# targeted view -- separate from (not a replacement for) the whole-document
# line-item scan below, which stays the primary extraction path since real
# section headers vary too much in exact wording to rely on alone.
_SECTION_PATTERNS = {
    'balance_sheet': re.compile(r'БАЛАНС|BALANCE\s+SHEET|АКТИВ(?:ЛАР)?\s+ВА\s+ПАССИВ', re.IGNORECASE),
    'p_and_l': re.compile(r'ДАРОМАД|ФОЙДА\s+(?:ВА|ЁКИ)\s+ЗАРАР|PROFIT\s+(?:AND|OR)\s+LOSS|INCOME\s+STATEMENT', re.IGNORECASE),
    'cash_flow': re.compile(r'ПУЛ\s+МАБЛАҒЛАРИ|CASH\s+FLOW', re.IGNORECASE),
    'notes': re.compile(r'ТУШУНТИРИШЛАР|NOTES\s+TO\s+FINANCIAL\s+STATEMENTS', re.IGNORECASE),
}


def _find_sections(text: str) -> dict:
    """
    Locates each recognizable section by its first keyword match, then
    spans that section's text up to the START of the next recognized
    section (in document order), not a fixed-size window -- a fixed window
    either cuts off long sections or wastes space on short ones.
    """
    hits = []
    for name, pattern in _SECTION_PATTERNS.items():
        match = pattern.search(text)
        if match:
            hits.append((match.start(), name))
    hits.sort()

    sections = {}
    for i, (start, name) in enumerate(hits):
        end = hits[i + 1][0] if i + 1 < len(hits) else len(text)
        sections[name] = text[start:end]
    return sections


def _find_company_name(text: str):
    """
    Finds a company name near a legal-entity suffix (e.g. LLC, MChJ). Uzbek
    company names are conventionally quoted (e.g. «Namuna» MChJ), so a
    quoted chunk on the same line as the suffix is preferred; otherwise
    falls back to the last few words before the suffix.
    """
    suffix_match = _COMPANY_SUFFIX_RE.search(text)
    if not suffix_match:
        return None

    line_start = text.rfind('\n', 0, suffix_match.start()) + 1
    window = text[line_start:suffix_match.end()]

    quoted_match = _QUOTED_NAME_RE.search(window)
    if quoted_match:
        return f'{quoted_match.group(1).strip()} {suffix_match.group()}'

    prefix = window[:suffix_match.start() - line_start].strip(' \'"«»()-—')
    words = prefix.split()
    name = ' '.join(words[-6:]) if words else ''
    return f'{name} {suffix_match.group()}'.strip()


def _parse_amount(raw: str):
    """
    Parses a numeric string into a float, tolerating the amount formats seen
    across real financial statements: comma or space thousands separators
    (including non-breaking spaces), a comma OR dot decimal separator
    (auto-detected), and parentheses or a leading minus for negatives.
    Returns None if the string isn't a parseable number.
    """
    s = raw.strip()
    if not s:
        return None

    negative = False
    if s.startswith('(') and s.endswith(')'):
        negative = True
        s = s[1:-1].strip()
    if s.startswith('-'):
        negative = True
        s = s[1:].strip()

    s = s.replace(' ', '').replace('\xa0', '')
    if not s:
        return None

    if ',' in s and '.' in s:
        # Whichever separator appears last is the decimal point.
        if s.rfind(',') > s.rfind('.'):
            s = s.replace('.', '').replace(',', '.')
        else:
            s = s.replace(',', '')
    elif ',' in s:
        parts = s.split(',')
        if len(parts) == 2 and len(parts[1]) in (1, 2):
            s = s.replace(',', '.')  # single comma, 1-2 trailing digits -> decimal
        else:
            s = s.replace(',', '')  # otherwise a thousands separator
    elif '.' in s:
        parts = s.split('.')
        if len(parts) > 2 or len(parts[-1]) == 3:
            s = s.replace('.', '')  # multiple dots, or 3 trailing digits -> thousands separator

    try:
        value = float(s)
    except ValueError:
        return None
    return -value if negative else value


def _clean_label(name_part: str) -> tuple:
    """
    Strips an optional leading numeric account code and, for bilingual
    "Uzbek label / English label" rows, keeps only the segment after the
    last "/". Returns (code, cleaned_label).
    """
    code = None
    code_match = _LEADING_CODE_RE.match(name_part)
    if code_match:
        code = code_match.group(1)
        name_part = name_part[code_match.end():].strip()

    if '/' in name_part:
        name_part = name_part.rsplit('/', 1)[-1].strip()

    return code, name_part


def _is_header_like(name_part: str) -> bool:
    """True for a "label" that's actually leftover digits/punctuation from a
    header/column-title row, not a real account name."""
    return not name_part or re.fullmatch(r'[\d\-\s.]*', name_part) is not None


def _extract_line_items(text: str, source: str) -> dict:
    """
    Scans text line by line for "<optional code> <name> <amount(s)>" rows,
    without requiring a rigid account-code prefix or a specific number
    format -- real statement text extracted from a PDF rarely preserves
    strict column alignment, decimals/space-thousands are common, and many
    line items have no visible code at all.

    Prefers a two-column match (current + prior period, common in
    comparative statements) and falls back to a single trailing amount.
    """
    line_items = {}
    for raw_line in text.split('\n'):
        line = raw_line.strip()
        if not line or _DATE_LINE_RE.search(line):
            continue

        two_col_match = _TWO_COLUMN_RE.match(line)
        if two_col_match:
            label_raw = two_col_match.group('label').strip()
            label_words = label_raw.split()
            # If the label's last word is itself all digits, this is almost
            # certainly a single space-grouped number ("30 000 000") that
            # got mis-split across the two "columns", not a real two-column
            # row -- a real account name essentially never ends in a bare
            # number. Fall through to single-amount handling instead.
            if label_words and re.fullmatch(r'[\d,.]+', label_words[-1]):
                two_col_match = None

        if two_col_match:
            code, name_part = _clean_label(label_raw)
            if _is_header_like(name_part):
                continue
            amount = _parse_amount(two_col_match.group('cur'))
            prior_amount = _parse_amount(two_col_match.group('prev'))
            if amount is None:
                continue
            entry = {'code': code, 'amount': amount, 'source': source}
            if prior_amount is not None:
                entry['prior_amount'] = prior_amount
            line_items[name_part] = entry
            continue

        amount_match = _AMOUNT_RE.search(line)
        if not amount_match:
            continue

        amount_str = amount_match.group().strip()
        if sum(c.isdigit() for c in amount_str) < 2:
            continue  # skip isolated single digits (e.g. footnote markers)

        code, name_part = _clean_label(line[:amount_match.start()].strip())
        if _is_header_like(name_part):
            continue

        amount = _parse_amount(amount_str)
        if amount is None:
            continue

        line_items[name_part] = {
            'code': code,
            'amount': amount,
            'source': source,
        }

    return line_items


def extract_from_pdf(pdf_path: str) -> dict:
    """
    Extracts financial statement data from PDF.

    Returns:
        dict: {
            "company_name": str,
            "period_end": str,
            "line_items": {...},  # Balance sheet + P&L lines
            "notes": {...}        # Notes to financial statements
        }
    """
    data = {"line_items": {}, "notes": {}}

    try:
        full_text = _pdfminer_extract_text(pdf_path) or ""

        # Simple extraction: look for balance sheet and P&L sections
        # (In production, use ML-based table detection; for MVP, regex is fine)

        # Extract company name (usually on first page)
        company_name = _find_company_name(full_text[:2000])
        if company_name:
            data['company_name'] = company_name

        # Extract reporting period (look for date pattern YYYY-MM-DD or Uzbek date)
        period_match = re.search(r'(\d{4})-(\d{2})-(\d{2})', full_text)
        if period_match:
            data['period_end'] = f"{period_match.group(1)}-{period_match.group(2)}-{period_match.group(3)}"

        data['line_items'] = _extract_line_items(full_text, source='PDF')

        # notes['raw'] is what gap_detector.py searches for keyword matches,
        # so it holds the FULL text (a substring search over it is cheap) --
        # not a truncated preview, which could miss real content in longer
        # documents.
        data['notes']['raw'] = full_text
        data['raw_text'] = full_text

        # Rough balance sheet / P&L / cash-flow / notes regions, each with
        # its own line-item scan. This is additional structure on top of
        # (not a replacement for) the whole-document scan above, since
        # section headers vary too much in exact wording to be the only
        # source of line items -- a document whose headers don't match
        # would otherwise yield nothing at all.
        data['sections'] = {
            name: {
                'text': section_text,
                'line_items': _extract_line_items(section_text, source='PDF'),
            }
            for name, section_text in _find_sections(full_text).items()
        }

    except Exception as e:
        data['error'] = str(e)

    return data


def extract_from_excel(excel_path: str) -> dict:
    """
    Extracts financial statement data from Excel (XLSX).

    Returns: same dict structure as PDF
    """
    data = {"line_items": {}, "notes": {}}

    wb = None
    try:
        wb = load_workbook(excel_path, read_only=True)

        # Assume first sheet is balance sheet, second is P&L
        for sheet_idx, sheet_name in enumerate(wb.sheetnames[:2]):
            ws = wb[sheet_name]

            for row in ws.iter_rows(min_row=1, max_row=200, values_only=True):
                if len(row) < 2 or not row[0]:
                    continue

                account_name = str(row[0]).strip()
                if not account_name:
                    continue

                # The amount isn't always in column B (e.g. current/prior
                # period columns) -- take the last cell in the row that
                # parses as a number.
                amount = None
                for cell in reversed(row[1:]):
                    if cell is None:
                        continue
                    amount = cell if isinstance(cell, (int, float)) else _parse_amount(str(cell))
                    if amount is not None:
                        break

                if amount is not None:
                    data['line_items'][account_name] = {
                        'sheet': sheet_name,
                        'amount': float(amount),
                        'source': 'Excel'
                    }

    except Exception as e:
        data['error'] = str(e)
    finally:
        # read_only workbooks keep the underlying file handle open until
        # closed; on Windows this locks the file and blocks deletion.
        if wb is not None:
            wb.close()

    return data


def extract_statement(file_path: str) -> dict:
    """
    Main entry point: detects file type and extracts data.
    """
    lower_path = file_path.lower()
    if lower_path.endswith('.pdf'):
        return extract_from_pdf(file_path)
    elif lower_path.endswith(('.xlsx', '.xls')):
        return extract_from_excel(file_path)
    else:
        return {'error': 'Unsupported file type'}
